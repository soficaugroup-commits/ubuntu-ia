"""Extraction de texte depuis Excel, CSV et PowerPoint."""

from __future__ import annotations

from pathlib import Path

from backend.ingestion.errors import IngestionError


def _decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise IngestionError("Le fichier n'a pas pu être décodé.")


def extract_csv(path: Path) -> str:
    import csv
    from io import StringIO

    text = _decode_bytes(path.read_bytes())
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[str] = []
    reader = csv.reader(StringIO(text), dialect)
    for row in reader:
        cells = [cell.strip() for cell in row]
        if any(cells):
            rows.append(" | ".join(cells))
    if not rows:
        raise IngestionError("Le fichier CSV ne contient aucune donnée exploitable.")
    return "\n".join(rows)


def _sheet_as_text(title: str, rows: list[list[str]]) -> str:
    lines = [f"# Feuille : {title}"]
    for row in rows:
        if any(cell.strip() for cell in row):
            lines.append(" | ".join(row))
    return "\n".join(lines)


def _xlsx_images(workbook) -> list[bytes]:
    blobs: list[bytes] = []
    for sheet in workbook.worksheets:
        for image in getattr(sheet, "_images", []) or []:
            try:
                data = image._data() if hasattr(image, "_data") else None
            except Exception:
                data = None
            if data:
                blobs.append(data)
    return blobs


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filename=str(path), data_only=True)
    except Exception as exc:
        raise IngestionError(f"Classeur Excel illisible : {exc}") from exc

    parts: list[str] = []
    image_blobs: list[bytes] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell).strip() for cell in row])
            block = _sheet_as_text(sheet.title, rows)
            if "\n" in block:
                parts.append(block)
        image_blobs = _xlsx_images(workbook)
    finally:
        workbook.close()

    if image_blobs:
        from backend.ingestion.extract_visual import describe_images

        visuals = describe_images(image_blobs)
        parts.extend(f"# Image du classeur\n{visual}" for visual in visuals)

    if not parts:
        raise IngestionError("Le classeur Excel ne contient pas de données exploitables.")
    return "\n\n".join(parts)


def extract_xls(path: Path) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise IngestionError(
            "La lecture des anciens fichiers .xls n'est pas installée."
        ) from exc

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise IngestionError(f"Ancien classeur Excel illisible : {exc}") from exc

    parts: list[str] = []
    for sheet in book.sheets():
        rows = [
            [str(sheet.cell_value(row_idx, col_idx)).strip() for col_idx in range(sheet.ncols)]
            for row_idx in range(sheet.nrows)
        ]
        block = _sheet_as_text(sheet.name, rows)
        if "\n" in block:
            parts.append(block)
    if not parts:
        raise IngestionError("L'ancien classeur Excel ne contient pas de données exploitables.")
    return "\n\n".join(parts)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise IngestionError(f"Présentation PowerPoint illisible : {exc}") from exc

    parts: list[str] = []
    image_blobs: list[bytes] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_lines = [f"# Diapositive {index}"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    slide_lines.append(text)
            if getattr(shape, "has_table", False):
                table_rows: list[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
                if table_rows:
                    slide_lines.append("\n".join(table_rows))
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image_blobs.append(shape.image.blob)
                except Exception:
                    pass
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            slide_lines.append(f"Notes : {notes}")
        if len(slide_lines) > 1:
            parts.append("\n".join(slide_lines))

    if image_blobs:
        from backend.ingestion.extract_visual import describe_images

        visuals = describe_images(image_blobs)
        parts.extend(f"# Image de la présentation\n{visual}" for visual in visuals)

    if not parts:
        raise IngestionError(
            "La présentation PowerPoint ne contient pas de texte ni d'image exploitable."
        )
    return "\n\n".join(parts)
