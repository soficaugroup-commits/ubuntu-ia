from pathlib import Path

from backend.ingestion.extract_file import extract_file
from backend.ingestion.validate import validate_file


def test_extract_csv(tmp_path: Path) -> None:
    path = tmp_path / "budget.csv"
    path.write_text("poste;montant\nLoyer;120000\n", encoding="utf-8")
    assert validate_file(path) == ".csv"
    text = extract_file(path)
    assert "Loyer" in text
    assert "120000" in text


def test_extract_xlsx(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "note.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "T1"
    sheet["A1"] = "Poste"
    sheet["B1"] = "Montant"
    sheet["A2"] = "Fournitures"
    sheet["B2"] = 45000
    book.save(path)

    assert validate_file(path) == ".xlsx"
    text = extract_file(path)
    assert "T1" in text
    assert "Fournitures" in text
    assert "45000" in text


def test_extract_pptx(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    path = tmp_path / "brief.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.paragraphs[0].text = "Procédure d'achat interne"
    box.text_frame.paragraphs[0].font.size = Pt(24)
    deck.save(path)

    assert validate_file(path) == ".pptx"
    text = extract_file(path)
    assert "Procédure d'achat interne" in text


def test_reject_fake_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "faux.xlsx"
    path.write_text("ceci n'est pas excel", encoding="utf-8")
    try:
        validate_file(path)
    except Exception as exc:
        assert "contenu réel" in str(exc)
    else:
        raise AssertionError("un faux xlsx doit être refusé")


def test_extract_docx_table(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Note de frais")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Poste"
    table.cell(0, 1).text = "Montant"
    table.cell(1, 0).text = "Déplacement"
    table.cell(1, 1).text = "25000"
    document.save(path)

    assert validate_file(path) == ".docx"
    text = extract_file(path)
    assert "Note de frais" in text
    assert "Déplacement" in text
    assert "25000" in text


def test_validate_xls_ole(tmp_path: Path) -> None:
    path = tmp_path / "ancien.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    assert validate_file(path) == ".xls"


def test_reject_legacy_doc(tmp_path: Path) -> None:
    path = tmp_path / "charte.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    try:
        validate_file(path)
    except Exception as exc:
        assert ".docx" in str(exc)
    else:
        raise AssertionError("un ancien .doc doit demander une conversion")


def test_validate_png_magic(tmp_path: Path) -> None:
    path = tmp_path / "pixel.png"
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        b"\x00\x00\x00\rIHDR"
        b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00"
        b"\x90wS\xde"
        b"\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05"
        b"\x18\xd8N"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    assert validate_file(path) == ".png"
