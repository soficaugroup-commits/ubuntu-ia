"""Extraction de texte depuis les formats indexables, y compris les fichiers anciens."""

from pathlib import Path

from backend.ingestion.errors import IngestionError
from backend.ingestion.extract_office import (
    extract_csv,
    extract_pptx,
    extract_xls,
    extract_xlsx,
)
from backend.ingestion.extract_visual import extract_image, extract_pdf_visual
from backend.ingestion.validate import validate_file


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise IngestionError(f"PDF illisible ou corrompu : {exc}") from exc

    text = "\n\n".join(part.strip() for part in pages if part.strip())
    if text:
        return text
    return extract_pdf_visual(path)


def _docx_images(document) -> list[bytes]:
    blobs: list[bytes] = []
    for rel in document.part.rels.values():
        if "image" not in getattr(rel, "reltype", ""):
            continue
        try:
            blobs.append(rel.target_part.blob)
        except Exception:
            continue
    return blobs


def extract_docx(path: Path) -> str:
    from docx import Document

    try:
        document = Document(str(path))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        image_blobs = _docx_images(document)
    except Exception as exc:
        raise IngestionError(f"Document Word illisible : {exc}") from exc

    parts = [part for part in parts if part]
    if image_blobs:
        from backend.ingestion.extract_visual import describe_images

        visuals = describe_images(image_blobs)
        parts.extend(f"# Image du document\n{visual}" for visual in visuals)

    if not parts:
        raise IngestionError("Le document Word ne contient pas de texte exploitable.")
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = raw.decode(encoding).strip()
            break
        except UnicodeDecodeError:
            text = ""
    if not text:
        raise IngestionError("Le fichier texte n'a pas pu être décodé.")
    return text


def extract_file(path: str | Path) -> str:
    """Valide le fichier puis retourne le texte brut exploitable par le modèle."""
    resolved = Path(path).expanduser().resolve()
    kind = validate_file(resolved)
    if kind == ".pdf":
        return extract_pdf(resolved)
    if kind == ".docx":
        return extract_docx(resolved)
    if kind in {".xlsx", ".xlsm"}:
        return extract_xlsx(resolved)
    if kind == ".xls":
        return extract_xls(resolved)
    if kind == ".csv":
        return extract_csv(resolved)
    if kind == ".pptx":
        return extract_pptx(resolved)
    if kind in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp"}:
        return extract_image(resolved)
    return extract_text(resolved)


if __name__ == "__main__":
    import argparse
    import sys

    parser = argparse.ArgumentParser(description="Extraire le texte d'un fichier.")
    parser.add_argument("fichier")
    args = parser.parse_args()
    try:
        print(extract_file(args.fichier))
    except IngestionError as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
